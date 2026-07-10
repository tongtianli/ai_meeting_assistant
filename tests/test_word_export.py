"""任务 5：Word 导出——docxtpl 实时渲染、重命名后真名即时生效、降级分支。"""
import io
import wave

from docx import Document
from fastapi.testclient import TestClient

from app.core.config import settings
from app.main import app
from tests.conftest import auth_headers, requires_db

pytestmark = requires_db

_DOCX_MT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _wav_bytes(seconds: float = 2.0, rate: int = 16000) -> bytes:
    buf = io.BytesIO()
    with wave.open(buf, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(rate)
        w.writeframes(b"\x00\x00" * int(rate * seconds))
    return buf.getvalue()


def _all_text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_export_docx_structured_with_rename(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    with TestClient(app) as client:
        headers = auth_headers(client)
        mid = client.post(
            "/api/meetings",
            files={"file": ("w.wav", _wav_bytes(), "audio/wav")},
            data={"title": "导出测试"},
            headers=headers,
        ).json()["id"]

        # 摘要完成后重命名 TODO 负责人（speaker_001 → 张三）
        resp = client.post(
            f"/api/meetings/{mid}/speaker-bindings",
            json={"speaker_label": "speaker_001", "name": "张三"},
            headers=headers,
        )
        assert resp.status_code == 201

        resp = client.get(f"/api/meetings/{mid}/export.docx", headers=headers)
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(_DOCX_MT)
        assert "attachment" in resp.headers["content-disposition"]

        text = _all_text(resp.content)
        assert "项目周会" in text  # mock 纪要标题
        assert "会议总结" in text
        assert "完成上传接口的开发" in text  # TODO 表格
        # 导出时通过来源 segment 现算负责人：重命名后立即用真名，无需重新摘要
        assert "张三" in text
        assert "00:00:0" in text  # TODO 来源时间戳
        # 模板标签不应泄漏到成品
        assert "{%" not in text and "{{" not in text


def test_export_docx_404_before_summary() -> None:
    with TestClient(app) as client:
        headers = auth_headers(client)
        resp = client.get(
            "/api/meetings/00000000-0000-0000-0000-0000000000ff/export.docx",
            headers=headers,
        )
        assert resp.status_code == 404


def test_export_docx_degraded(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    import app.services.summarize as summarize_mod
    from app.services.llm.base import LLMProvider, LLMResponse
    from app.services.llm.router import LLMRouter

    class _Broken(LLMProvider):
        name = "broken"
        model = "broken-model"

        async def complete(self, system, user, json_mode=True, temperature=0.2):
            text = "不是 JSON" if json_mode else "纯文本纪要：讨论了项目进展。"
            return LLMResponse(text=text, provider=self.name, model=self.model)

    monkeypatch.setattr(
        summarize_mod, "build_router", lambda: LLMRouter([_Broken()])
    )
    with TestClient(app) as client:
        headers = auth_headers(client)
        mid = client.post(
            "/api/meetings",
            files={"file": ("d.wav", _wav_bytes(), "audio/wav")},
            data={"title": "降级导出"},
            headers=headers,
        ).json()["id"]

        resp = client.get(f"/api/meetings/{mid}/export.docx", headers=headers)
        assert resp.status_code == 200
        text = _all_text(resp.content)
        assert "纯文本纪要" in text
        assert "{%" not in text
