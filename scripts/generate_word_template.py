"""生成公司纪要 Word 模板（docxtpl/jinja2 标签）。

版式复刻公司既有纪要（表格式公文）：
四列表格抬头（主题/时间/地点/主持/参会/记录人/重要程度）
+「会议主要内容：」正文按议题分组（一、议题（责任人：X）+ 编号条目）
+ 参会人员签字、备注留空供打印手写。
用法：uv run python scripts/generate_word_template.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parent.parent / "app" / "templates" / "summary_template.docx"

SONG = "宋体"


def _song(run, size: int = 12, bold: bool = False) -> None:
    run.font.name = SONG
    run.font.size = Pt(size)
    run.bold = bold
    # 中文字体需额外设置 eastAsia，否则 Word 里中文回退默认字体
    run._element.rPr.rFonts.set(qn("w:eastAsia"), SONG)


def _label(cell, text: str) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _song(p.add_run(text), bold=True)


def _value(cell, text: str, center: bool = True) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _song(p.add_run(text))


def main() -> None:
    doc = Document()
    # 标题行
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _song(heading.add_run("会 议 纪 要"), size=18, bold=True)

    table = doc.add_table(rows=7, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = (Cm(2.6), Cm(5.6), Cm(2.6), Cm(5.6))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w

    # 第 1 行：会议主题（值格横向合并）
    _label(table.rows[0].cells[0], "会议主题")
    merged = table.rows[0].cells[1].merge(table.rows[0].cells[3])
    _value(merged, "{{ title }}")

    # 第 2 行：会议时间 | 会议地点
    _label(table.rows[1].cells[0], "会议时间")
    _value(table.rows[1].cells[1], "{{ meeting_time }}")
    _label(table.rows[1].cells[2], "会议地点")
    _value(table.rows[1].cells[3], "{{ location }}")

    # 第 3 行：会议主持 | 参会人员
    _label(table.rows[2].cells[0], "会议主持")
    _value(table.rows[2].cells[1], "{{ host }}")
    _label(table.rows[2].cells[2], "参会人员")
    _value(table.rows[2].cells[3], "{{ participants }}")

    # 第 4 行：记录人 | 重要程度（未指定时三档并列，打印后圈选）
    _label(table.rows[3].cells[0], "记录人")
    _value(table.rows[3].cells[1], "{{ recorder }}")
    _label(table.rows[3].cells[2], "重要程度")
    _value(table.rows[3].cells[3], "{{ importance_line }}")

    # 第 5 行：会议议题（正文大格）
    _label(table.rows[4].cells[0], "会议议题")
    body = table.rows[4].cells[1].merge(table.rows[4].cells[3])
    body.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    _song(body.paragraphs[0].add_run("会议主要内容："), bold=True)
    # 降级（纯文本）分支
    body.add_paragraph("{%p if degraded %}")
    _song(body.add_paragraph().add_run("{{ plain_text }}"))
    body.add_paragraph("{%p endif %}")
    # 结构化分支：议题分组
    body.add_paragraph("{%p if not degraded %}")
    body.add_paragraph("{%p for t in topics %}")
    _song(
        body.add_paragraph().add_run("{{ t.num }}、{{ t.title }}{{ t.owner_suffix }}"),
        bold=True,
    )
    body.add_paragraph("{%p for it in t.entries %}")
    _song(body.add_paragraph().add_run("{{ loop.index }}. {{ it }}"))
    body.add_paragraph("{%p endfor %}")
    body.add_paragraph("{%p endfor %}")
    body.add_paragraph("{%p endif %}")

    # 第 6 行：参会人员签字（留空，打印手签）
    _label(table.rows[5].cells[0], "参会人员\n签字")
    sign = table.rows[5].cells[1].merge(table.rows[5].cells[3])
    _value(sign, "", center=False)

    # 第 7 行：备注
    _label(table.rows[6].cells[0], "备注")
    note = table.rows[6].cells[1].merge(table.rows[6].cells[3])
    _value(note, "", center=False)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"template written: {OUT}")


if __name__ == "__main__":
    main()
